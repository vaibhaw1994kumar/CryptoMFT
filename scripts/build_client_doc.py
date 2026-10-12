from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CLIENT_POC_OVERVIEW.md"
OUTPUT = ROOT / "CLIENT_POC_OVERVIEW.docx"


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text


def shade(cell_or_paragraph, fill: str) -> None:
    target = cell_or_paragraph._tc if hasattr(cell_or_paragraph, "_tc") else cell_or_paragraph._p
    properties_name = "w:tcPr" if hasattr(cell_or_paragraph, "_tc") else "w:pPr"
    properties = target.find(qn(properties_name))
    if properties is None:
        properties = OxmlElement(properties_name)
        target.insert(0, properties)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_code(document: Document, lines: list[str], language: str) -> None:
    if language == "mermaid":
        caption = document.add_paragraph()
        caption.add_run(
            "Diagram definition (rendered visually in CLIENT_POC_OVERVIEW.md)"
        ).italic = True
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["No Spacing"]
    shade(paragraph, "F1F5F9")
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = Pt(8)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = clean_inline(value)
        shade(cell, "1E3A8A")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = clean_inline(value)
    document.add_paragraph()


def build(source: Path = SOURCE, output: Path = OUTPUT) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    for style_name, size, color in (
        ("Title", 23, RGBColor(23, 37, 84)),
        ("Heading 1", 16, RGBColor(30, 58, 138)),
        ("Heading 2", 13, RGBColor(30, 64, 175)),
        ("Heading 3", 11, RGBColor(29, 78, 216)),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color

    lines = source.read_text(encoding="utf-8").splitlines()
    if lines and lines[0] == "---":
        end = lines.index("---", 1)
        lines = lines[end + 1 :]

    in_code = False
    code_language = ""
    code_lines: list[str] = []
    paragraph_lines: list[str] = []
    index = 0

    def flush_paragraph() -> None:
        if paragraph_lines:
            document.add_paragraph(clean_inline(" ".join(paragraph_lines)))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code(document, code_lines, code_language)
                code_lines.clear()
                code_language = ""
                in_code = False
            else:
                code_language = line[3:].strip()
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip() or line.startswith("<!--"):
            flush_paragraph()
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1]
            if separator.startswith("|") and "---" in separator:
                flush_paragraph()
                table_rows = [
                    [part.strip() for part in line.strip("|").split("|")]
                ]
                index += 2
                while index < len(lines) and lines[index].startswith("|"):
                    table_rows.append(
                        [
                            part.strip()
                            for part in lines[index].strip("|").split("|")
                        ]
                    )
                    index += 1
                add_table(document, table_rows)
                continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1 and not document.paragraphs:
                paragraph = document.add_paragraph(text, style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_heading(text, level=level)
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            document.add_paragraph(
                clean_inline(bullet.group(1)),
                style="List Bullet",
            )
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            document.add_paragraph(
                clean_inline(numbered.group(1)),
                style="List Number",
            )
            index += 1
            continue
        paragraph_lines.append(line.strip())
        index += 1

    flush_paragraph()

    footer = section.footer.paragraphs[0]
    footer.text = "Binance Trading Platform Proof of Concept - Client Guide"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    document.core_properties.title = "HFT Trading Platform Proof of Concept"
    document.core_properties.subject = (
        "Client scope, architecture, instructions, and acceptance criteria"
    )
    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    arguments = parser.parse_args()
    build(arguments.source, arguments.output)
